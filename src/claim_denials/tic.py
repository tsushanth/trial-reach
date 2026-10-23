"""Loader and analysis for the marketplace transparency workbook (split in-network / out-of-network format).

The workbook path comes from the caller (an environment variable in the scripts). No data and no data location lives in this repository.
Values `**`, `*`, `N/A`, blanks and non-numeric text are missing and are never imputed.
"""

from __future__ import annotations

import math
import statistics
from dataclasses import dataclass, field

MIN_CLAIMS = 1000
MIN_APPEALS = 10

COLS = {
    "state": "State", "issuer_id": "Issuer_ID", "issuer_name": "Issuer_Name",
    "rec_in": "Issuer_Claims_Received_In_Network", "rec_out": "Issuer_Claims_Received_Out_of_Network",
    "den_in": "Issuer_Claims_Denied_In_Network", "den_out": "Issuer_Claims_Denied_Out_of_Network",
    "resub_in": "Issuer_Claims_Resubmitted_In_Network", "resub_out": "Issuer_Claims_Resubmitted_Out_of_Network",
    "int_filed": "Issuer_Internal_Appeals_Filed", "int_over": "Issuer_Number_Internal_Appeals_Overturned",
    "ext_filed": "Issuer_External_Appeals_Filed", "ext_over": "Issuer_Number_External_Appeals_Overturned",
}
PLAN_REASONS = {
    "referral": "Plan_Number_Claims_Denied_Referral_Required", "out_of_network": "Plan_Number_Claims_Denied_Due_To_Out_Of_Network",
    "excluded": "Plan_Number_Claims_Denied_Services_Excluded", "benefit_limit": "Plan_Number_Claims_Denied_Due_To_Enrolle_Benefit_Limit_Reached",
    "not_covered": "Plan_Number_Claims_Denied_Due_To_Member_Not_Covered", "administrative": "Plan_Number_Claims_Denied_Due_To_Administrative_Reason",
    "other": "Plan_Number_Claims_Denied_Other",
}
PLAN_REASON_PREFIXES = {  # headers that are long or vary slightly across years are matched by prefix
    "medical_necessity": "Plan_Number_Claims_Denied_Not_Medically_Necessary",
    "investigational": "Plan_Number_Claims_Denied_Due_To_Investigational",
}


def num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return None if (isinstance(v, float) and math.isnan(v)) else float(v)
    s = str(v).replace(",", "").strip()
    try:
        return float(s)
    except ValueError:
        return None


@dataclass
class Issuer:
    state: str
    issuer_id: str
    name: str
    vals: dict = field(default_factory=dict)
    reasons: dict = field(default_factory=dict)

    @property
    def key(self):
        return (self.state, str(self.issuer_id))


def read_sheet(rows: list[tuple]) -> list[Issuer]:
    """rows: all rows of the sheet. Finds the header row by the Issuer_Name column. One Issuer per (state, issuer id)."""
    hi = next((i for i, r in enumerate(rows) if r and "Issuer_Name" in [str(c) for c in r]), None)
    if hi is None:
        raise ValueError("no header row with Issuer_Name")
    head = [str(c) if c is not None else "" for c in rows[hi]]
    pos = {h: i for i, h in enumerate(head)}
    out: dict = {}
    for r in rows[hi + 1:]:
        iid = r[pos["Issuer_ID"]] if "Issuer_ID" in pos else None
        if iid in (None, ""):
            continue
        key = (str(r[pos["State"]]), str(iid))
        if key not in out:
            it = Issuer(key[0], key[1], str(r[pos["Issuer_Name"]]))
            it.vals = {k: num(r[pos[c]]) if c in pos else None for k, c in COLS.items() if k not in ("state", "issuer_id", "issuer_name")}
            out[key] = it
        it = out[key]
        for k, c in PLAN_REASONS.items():
            if c in pos:
                v = num(r[pos[c]])
                if v is not None:
                    it.reasons[k] = it.reasons.get(k, 0.0) + v
        for k, pre in PLAN_REASON_PREFIXES.items():
            for h, i in pos.items():
                if h.startswith(pre):
                    v = num(r[i])
                    if v is not None:
                        it.reasons[k] = it.reasons.get(k, 0.0) + v
    return list(out.values())


def load_workbook_issuers(path: str, sheet_hint: str = "Ind QHP") -> list[Issuer]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    name = next((n for n in wb.sheetnames if sheet_hint in n), None)
    if name is None:
        raise ValueError(f"no sheet containing {sheet_hint!r}")
    return read_sheet(list(wb[name].iter_rows(values_only=True)))


def denial_rate(it: Issuer, net: str = "in"):
    rec, den = it.vals.get(f"rec_{net}"), it.vals.get(f"den_{net}")
    if rec is None or den is None or rec < MIN_CLAIMS:
        return None
    return den / rec


def plausible(it: Issuer) -> bool:
    v = it.vals
    for rec, den in (("rec_in", "den_in"), ("rec_out", "den_out")):
        if v.get(rec) is not None and v.get(den) is not None and v[den] > v[rec]:
            return False
    for filed, over in (("int_filed", "int_over"), ("ext_filed", "ext_over")):
        if v.get(filed) is not None and v.get(over) is not None and v[over] > v[filed]:
            return False
    return True


def quartiles(xs):
    xs = sorted(xs)
    if len(xs) < 4:
        return None
    q = statistics.quantiles(xs, n=4, method="inclusive")
    return {"p25": q[0], "median": q[1], "p75": q[2]}


def spearman(pairs):
    if len(pairs) < 3:
        return None

    def ranks(v):
        order = sorted(range(len(v)), key=lambda i: v[i])
        r = [0.0] * len(v)
        i = 0
        while i < len(order):
            j = i
            while j + 1 < len(order) and v[order[j + 1]] == v[order[i]]:
                j += 1
            for k in range(i, j + 1):
                r[order[k]] = (i + j) / 2 + 1
            i = j + 1
        return r

    a, b = ranks([p[0] for p in pairs]), ranks([p[1] for p in pairs])
    ma, mb = statistics.fmean(a), statistics.fmean(b)
    num_ = sum((x - ma) * (y - mb) for x, y in zip(a, b))
    den = math.sqrt(sum((x - ma) ** 2 for x in a) * sum((y - mb) ** 2 for y in b))
    return num_ / den if den else None


def year_summary(issuers: list[Issuer]) -> dict:
    ok = [i for i in issuers if plausible(i)]
    rates = [r for r in (denial_rate(i) for i in ok) if r is not None]
    summ = {"issuers": len(issuers), "implausible": len(issuers) - len(ok), "with_rate": len(rates)}
    if rates:
        summ.update(quartiles(rates) or {})
        summ["share_above_30pct"] = sum(r > 0.30 for r in rates) / len(rates)
        pr = sum(i.vals["den_in"] for i in ok if denial_rate(i) is not None) / sum(i.vals["rec_in"] for i in ok if denial_rate(i) is not None)
        summ["pooled_rate_claim_weighted"] = pr
    # appeals
    den = filed = over = ext_f = ext_o = 0.0
    n_int = n_ext = 0
    for i in ok:
        v = i.vals
        if None not in (v.get("den_in"), v.get("den_out"), v.get("int_filed")) and v["int_filed"] >= 0:
            den += v["den_in"] + v["den_out"]
            filed += v["int_filed"]
        if v.get("int_filed") is not None and v["int_filed"] >= MIN_APPEALS and v.get("int_over") is not None:
            over += v["int_over"]
            n_int += 1
            summ.setdefault("_int_pairs", []).append((v["int_over"], v["int_filed"]))
        if v.get("ext_filed") is not None and v["ext_filed"] >= MIN_APPEALS and v.get("ext_over") is not None:
            ext_f += v["ext_filed"]
            ext_o += v["ext_over"]
            n_ext += 1
    pairs = summ.pop("_int_pairs", [])
    if den:
        summ["internal_appeals_per_denied_claim"] = filed / den
    if pairs:
        summ["internal_overturn_pooled"] = sum(p[0] for p in pairs) / sum(p[1] for p in pairs)
        summ["issuers_with_internal_overturn_rate"] = n_int
    if ext_f:
        summ["external_overturn_pooled"] = ext_o / ext_f
        summ["issuers_with_external_overturn_rate"] = n_ext
    return summ


def persistence(prev: list[Issuer], cur: list[Issuer]) -> dict:
    p = {i.key: denial_rate(i) for i in prev if plausible(i)}
    c = {i.key: denial_rate(i) for i in cur if plausible(i)}
    both = [(p[k], c[k]) for k in p if k in c and p[k] is not None and c[k] is not None]
    out = {"issuers_in_both": len(both)}
    if both:
        out["spearman"] = spearman(both)
        out["share_changed_over_10pp"] = sum(abs(a - b) > 0.10 for a, b in both) / len(both)
    return out
