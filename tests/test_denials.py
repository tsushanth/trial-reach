import openpyxl
import pytest

from claim_denials import tic

HEAD = ["State", "Issuer_Name", "Issuer_ID", "Plan_ID", "Issuer_Claims_Received_In_Network", "Issuer_Claims_Received_Out_of_Network",
        "Issuer_Claims_Denied_In_Network", "Issuer_Claims_Denied_Out_of_Network", "Issuer_Internal_Appeals_Filed",
        "Issuer_Number_Internal_Appeals_Overturned", "Issuer_External_Appeals_Filed", "Issuer_Number_External_Appeals_Overturned",
        "Plan_Number_Claims_Denied_Due_To_Administrative_Reason"]


def sheet(rows):
    return [("Title",), ("Legend",), tuple(HEAD)] + rows


def row(state, name, iid, plan, rin, rout, din, dout, fi, oi, fe, oe, admin="5"):
    return (state, name, iid, plan, rin, rout, din, dout, fi, oi, fe, oe, admin)


def test_reader_dedupes_issuer_level_values_and_sums_plan_reasons():
    rows = sheet([row("TX", "Alpha", "1", "p1", 10000, 2000, 2000, 1000, 100, 40, 20, 5, "5"),
                  row("TX", "Alpha", "1", "p2", 10000, 2000, 2000, 1000, 100, 40, 20, 5, "7")])
    its = tic.read_sheet(rows)
    assert len(its) == 1 and its[0].vals["rec_in"] == 10000 and its[0].reasons["administrative"] == 12


def test_missing_markers_are_none_and_small_cells_have_no_rate():
    rows = sheet([row("TX", "A", "1", "p", "**", "*", "N/A", None, 0, 0, 0, 0), row("TX", "B", "2", "p", 500, 0, 100, 0, 0, 0, 0, 0)])
    a, b = tic.read_sheet(rows)
    assert a.vals["rec_in"] is None and tic.denial_rate(a) is None and tic.denial_rate(b) is None  # below MIN_CLAIMS
    assert tic.num("1,234") == 1234.0 and tic.num("**") is None and tic.num(True) is None


def test_plausibility_and_year_summary():
    good = [row("TX", f"I{k}", str(k), "p", 10000, 1000, 1000 + 300 * k, 100, 200, 100, 20, 8) for k in range(5)]
    bad = [row("TX", "Bad", "99", "p", 10000, 1000, 20000, 100, 10, 50, 0, 0)]  # denied > received, overturned > filed
    its = tic.read_sheet(sheet(good + bad))
    assert [tic.plausible(i) for i in its] == [True] * 5 + [False]
    s = tic.year_summary(its)
    assert s["implausible"] == 1 and s["with_rate"] == 5 and 0.1 < s["median"] < 0.3
    assert s["internal_overturn_pooled"] == pytest.approx(0.5) and s["external_overturn_pooled"] == pytest.approx(0.4)


def test_spearman_and_persistence():
    assert tic.spearman([(1, 1), (2, 2), (3, 3), (4, 4)]) == pytest.approx(1.0)
    assert tic.spearman([(1, 4), (2, 3), (3, 2), (4, 1)]) == pytest.approx(-1.0)
    assert tic.spearman([(1, 1)]) is None
    y1 = tic.read_sheet(sheet([row("TX", f"I{k}", str(k), "p", 10000, 0, 1000 * (k + 1), 0, 0, 0, 0, 0) for k in range(4)]))
    y2 = tic.read_sheet(sheet([row("TX", f"I{k}", str(k), "p", 10000, 0, 1100 * (k + 1), 0, 0, 0, 0, 0) for k in range(4)]))
    r = tic.persistence(y1, y2)
    assert r["issuers_in_both"] == 4 and r["spearman"] == pytest.approx(1.0) and r["share_changed_over_10pp"] == 0.0


def test_workbook_loader_reads_the_right_sheet(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Disclaimer"
    ws = wb.create_sheet("Transparency 2099 - Ind QHP")
    for r in sheet([row("TX", "Alpha", "1", "p", 10000, 0, 1000, 0, 0, 0, 0, 0)]):
        ws.append(r)
    p = tmp_path / "x.xlsx"
    wb.save(p)
    assert [i.name for i in tic.load_workbook_issuers(str(p))] == ["Alpha"]
